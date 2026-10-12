"""把 PDF / Word / Excel / PowerPoint 转成纯文本，供双AI对话台把附件内容放进提问里。

用法：python extract_text.py <文件> <输出目录>
输出：一行 JSON {"text": "...", "pages": N, "images": ["page-1.png", ...], "engine": "..."}
扫描版 PDF（几乎没有文字）会把前 4 页渲染成 PNG 放进输出目录，交给模型「看图」。
缺库时尽量用标准库兜底（docx/pptx 本质是 zip+XML）。
"""
import json
import os
import re
import sys
import zipfile

MAX_ROWS_PER_SHEET = 2000
SCAN_PAGES = 4
SCAN_THRESHOLD = 50


def pdf_text(path, outdir):
    try:
        import fitz  # PyMuPDF
    except ImportError:
        fitz = None
    if fitz is not None:
        doc = fitz.open(path)
        parts = [f"[p.{i + 1}]\n{page.get_text().strip()}" for i, page in enumerate(doc)]
        raw = "".join(page.get_text() for page in doc)
        images = []
        if len(raw.strip()) < SCAN_THRESHOLD:
            for i, page in enumerate(doc):
                if i >= SCAN_PAGES:
                    break
                name = f"page-{i + 1}.png"
                page.get_pixmap(dpi=110).save(os.path.join(outdir, name))
                images.append(name)
        return {"text": "\n\n".join(parts) if raw.strip() else "", "pages": doc.page_count, "images": images, "engine": "pymupdf"}
    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = [f"[p.{i + 1}]\n{(page.extract_text() or '').strip()}" for i, page in enumerate(reader.pages)]
    text = "\n\n".join(parts)
    return {"text": text if re.sub(r"\[p\.\d+\]|\s", "", text) else "", "pages": len(reader.pages), "images": [], "engine": "pypdf"}


def xml_texts(data, tag):
    return [m.group(1) for m in re.finditer(rf"<{tag}[^>]*>([^<]*)</{tag}>", data)]


def docx_text(path, _outdir):
    try:
        import docx
        d = docx.Document(path)
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                lines.append("\t".join(cell.text.strip() for cell in row.cells))
        return {"text": "\n".join(lines), "pages": 0, "images": [], "engine": "python-docx"}
    except ImportError:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8", "replace")
        paras = [" ".join(xml_texts(p, "w:t")) for p in xml.split("</w:p>")]
        return {"text": "\n".join(p for p in paras if p.strip()), "pages": 0, "images": [], "engine": "zip"}


def xlsx_text(path, _outdir):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS_PER_SHEET:
                out.append(f"…（{ws.title} 共超过 {MAX_ROWS_PER_SHEET} 行，后面省略）")
                break
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                out.append("\t".join(cells).rstrip())
    return {"text": "\n".join(out), "pages": len(wb.worksheets), "images": [], "engine": "openpyxl"}


def pptx_text(path, _outdir):
    with zipfile.ZipFile(path) as z:
        slides = sorted(
            (n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
            key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)),
        )
        parts = []
        for i, name in enumerate(slides):
            xml = z.read(name).decode("utf-8", "replace")
            parts.append(f"[slide {i + 1}]\n" + "\n".join(t for t in xml_texts(xml, "a:t") if t.strip()))
    return {"text": "\n\n".join(parts), "pages": len(slides), "images": [], "engine": "zip"}


HANDLERS = {".pdf": pdf_text, ".docx": docx_text, ".xlsx": xlsx_text, ".xlsm": xlsx_text, ".pptx": pptx_text}


def main():
    if len(sys.argv) != 3:
        print(json.dumps({"error": "usage: extract_text.py <file> <outdir>"}))
        return 2
    path, outdir = sys.argv[1], sys.argv[2]
    handler = HANDLERS.get(os.path.splitext(path)[1].lower())
    if handler is None:
        print(json.dumps({"error": "unsupported"}))
        return 2
    try:
        result = handler(path, outdir)
    except Exception as exc:  # 交给 Node 端记录，不让异常栈刷屏
        print(json.dumps({"error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
